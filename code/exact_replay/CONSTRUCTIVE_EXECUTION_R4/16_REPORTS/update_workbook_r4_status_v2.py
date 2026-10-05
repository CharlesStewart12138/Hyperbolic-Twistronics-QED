"""Compatibility-repaired versioned workbook updater (v2)."""
from __future__ import annotations

import copy
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties


R4 = Path(__file__).resolve().parents[1]
SOURCE = R4 / "00_FROZEN_INPUTS/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan.xlsx"
ROWS = R4 / "16_REPORTS/WORKBOOK_CONSTRUCTIVE_TASK_ROWS.tsv"
OUTPUT = R4 / "16_REPORTS/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan_R4_Status.xlsx"
AUDIT = R4 / "16_REPORTS/WORKBOOK_R4_STATUS_UPDATE_AUDIT.json"
SHEET = "Master Tasks"
ORIGINAL_LAST_ROW = 261


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def value_fingerprint(ws, last_row: int) -> str:
    payload = [[ws.cell(r, c).value for c in range(1, 19)] for r in range(1, last_row + 1)]
    encoded = json.dumps(payload, ensure_ascii=False, default=str, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest().upper()


source_hash_before = sha256(SOURCE)
source_wb = load_workbook(SOURCE, read_only=False, data_only=False)
source_ws = source_wb[SHEET]
assert source_ws.max_row == ORIGINAL_LAST_ROW
headers = [source_ws.cell(1, c).value for c in range(1, 19)]
source_prefix_hash = value_fingerprint(source_ws, ORIGINAL_LAST_ROW)

with ROWS.open("r", encoding="utf-8", newline="") as handle:
    proposed = list(csv.DictReader(handle, delimiter="\t"))
assert proposed and list(proposed[0]) == headers

output_wb = load_workbook(SOURCE, read_only=False, data_only=False)
output_ws = output_wb[SHEET]
for offset, record in enumerate(proposed, start=1):
    target_row = ORIGINAL_LAST_ROW + offset
    output_ws.row_dimensions[target_row].height = output_ws.row_dimensions[ORIGINAL_LAST_ROW].height
    for column, header in enumerate(headers, start=1):
        source_cell = output_ws.cell(ORIGINAL_LAST_ROW, column)
        target_cell = output_ws.cell(target_row, column, record[header])
        if source_cell.has_style:
            target_cell._style = copy.copy(source_cell._style)
        target_cell.number_format = source_cell.number_format
        target_cell.font = copy.copy(source_cell.font)
        target_cell.fill = copy.copy(source_cell.fill)
        target_cell.border = copy.copy(source_cell.border)
        target_cell.alignment = copy.copy(source_cell.alignment)
        target_cell.protection = copy.copy(source_cell.protection)

new_last_row = ORIGINAL_LAST_ROW + len(proposed)
output_ws.tables["MasterTasksTable"].ref = f"A1:R{new_last_row}"
if getattr(output_wb, "calculation", None) is None:
    output_wb.calculation = CalcProperties()
output_wb.calculation.fullCalcOnLoad = True
output_wb.calculation.forceFullCalc = True
output_wb.calculation.calcMode = "auto"
output_wb.save(OUTPUT)

reloaded = load_workbook(OUTPUT, read_only=False, data_only=False)
reloaded_ws = reloaded[SHEET]
source_hash_after = sha256(SOURCE)
assert source_hash_after == source_hash_before
assert value_fingerprint(reloaded_ws, ORIGINAL_LAST_ROW) == source_prefix_hash
assert reloaded_ws.max_row == new_last_row
assert reloaded_ws.tables["MasterTasksTable"].ref == f"A1:R{new_last_row}"

expected_ids = [row["Task ID"] for row in proposed]
actual_ids = [reloaded_ws.cell(ORIGINAL_LAST_ROW + i, 1).value for i in range(1, len(proposed) + 1)]
assert actual_ids == expected_ids

formula_count_source = sum(
    1 for ws in source_wb.worksheets for row in ws.iter_rows() for cell in row
    if isinstance(cell.value, str) and cell.value.startswith("=")
)
formula_count_output = sum(
    1 for ws in reloaded.worksheets for row in ws.iter_rows() for cell in row
    if isinstance(cell.value, str) and cell.value.startswith("=")
)
assert formula_count_output == formula_count_source

audit = {
    "schema_version": "1.0",
    "classification": "PASS_VERSIONED_COPY",
    "source_path": "00_FROZEN_INPUTS/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan.xlsx",
    "source_sha256_before": source_hash_before,
    "source_sha256_after": source_hash_after,
    "source_mutated": False,
    "output_path": "16_REPORTS/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan_R4_Status.xlsx",
    "output_sha256": sha256(OUTPUT),
    "sheet": SHEET,
    "original_last_row": ORIGINAL_LAST_ROW,
    "rows_appended": len(proposed),
    "new_last_row": new_last_row,
    "appended_task_ids": actual_ids,
    "original_cell_value_fingerprint": source_prefix_hash,
    "output_original_region_fingerprint": value_fingerprint(reloaded_ws, ORIGINAL_LAST_ROW),
    "original_region_preserved": True,
    "table_range": reloaded_ws.tables["MasterTasksTable"].ref,
    "formula_count_source": formula_count_source,
    "formula_count_output": formula_count_output,
    "recalculate_on_open": True,
    "supersedes": "16_REPORTS/update_workbook_r4_status.py",
}
AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({"classification": audit["classification"], "output": str(OUTPUT), "rows_appended": len(proposed), "output_sha256": audit["output_sha256"]}, indent=2))
