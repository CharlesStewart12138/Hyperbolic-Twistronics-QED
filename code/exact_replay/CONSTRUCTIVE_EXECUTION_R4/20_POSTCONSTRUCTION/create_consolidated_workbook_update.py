"""Create the single consolidated post-construction workbook revision."""
from __future__ import annotations

import copy
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.workbook.properties import CalcProperties


R4 = Path(__file__).resolve().parents[1]
SOURCE = R4 / "00_FROZEN_INPUTS/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan.xlsx"
PREVIEW = R4 / "19_RESOURCE_CLOSURE/WORKBOOK_UPDATE_PREVIEW.tsv"
OUTPUT = R4 / "20_POSTCONSTRUCTION/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan_POST_CONSTRUCTION_FINAL.xlsx"
AUDIT = R4 / "20_POSTCONSTRUCTION/AUTHORITATIVE_WORKBOOK_CONSOLIDATED_UPDATE_AUDIT.json"
H_MATH = json.loads((R4 / "17_FINAL_FREEZE/CAND-R4-0005_ROOT_CERTIFICATE.json").read_text(encoding="utf-8"))["H_MATH"]
ORIGINAL_LAST_ROW = 261


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def fingerprint(ws, last_row: int) -> str:
    values = [[ws.cell(r, c).value for c in range(1, 19)] for r in range(1, last_row + 1)]
    return hashlib.sha256(json.dumps(values, ensure_ascii=False, default=str, separators=(",", ":")).encode("utf-8")).hexdigest().upper()


source_hash_before = sha256(SOURCE)
source_wb = load_workbook(SOURCE, data_only=False, read_only=False)
source_ws = source_wb["Master Tasks"]
headers = [source_ws.cell(1, c).value for c in range(1, 19)]
source_prefix = fingerprint(source_ws, ORIGINAL_LAST_ROW)

task_specs = [
    ("PF-GRP-CONSTRUCTIVE-R4", "Construct and certify production A1 quotient", "Blocked", "Mathematical construction closed; resource specification remains incomplete."),
    ("PF-GRP-CONSTRUCTIVE-R4-FROZEN", "Freeze exact constructive inputs and H_MATH", "Done", H_MATH),
    ("PF-GRP-CONSTRUCTIVE-R4-ARITH", "Build arithmetic separator factors", "Done", "p=3 and p=5 exact factors."),
    ("PF-GRP-CONSTRUCTIVE-R4-SEPLIB", "Build witness/separator library", "Done", "14 C8 witness orbits; five standalone separator columns."),
    ("PF-GRP-CONSTRUCTIVE-R4-CEGAR", "Run kernel-first CEGAR sequence", "Done", "Five sequential candidates; CAND-R4-0005 global PASS."),
    ("PF-GRP-CONSTRUCTIVE-R4-PQUOT", "Audit C8-stable class-2 p=2 branches", "Done", "All 7 D1 and 19 D2 stable spaces audited."),
    ("PF-GRP-CONSTRUCTIVE-R4-PRODUCT", "Form/minimize actual subdirect product", "Done", "Q=SL(2,9)xC4xC4xC2xC2, order 46080; factor-irredundant."),
    ("PF-GRP-CONSTRUCTIVE-R4-GLOBAL", "Certify strict global systole", "Done", "Two complete 785639753-record exact scans; sys/a_B>6."),
    ("PF-GRP-CONSTRUCTIVE-R4-PERM", "Compute exact permutation degrees", "Done", "mu(Q)=92; mu_tr(Q)=5120."),
    ("PF-GRP-CONSTRUCTIVE-R4-Q24", "Resolve historical q24 relation", "Done", "Branch A: Q lies outside the faithful transitive degree-24 domain."),
    ("PF-GRP-CONSTRUCTIVE-R4-RESOURCE", "Complete numerical production resource gate", "Blocked", "Resource UNVERIFIED: 22 missing project-decision rows and one implementation-dependent support count."),
    ("PF-GRP-CONSTRUCTIVE-R4-A1", "Freeze LEVEL_A1_CONSTRUCTIVE_R4", "Blocked", "Not frozen; Resource Gate is not PASS."),
    ("PF-GRP-CONSTRUCTIVE-R4-PAPER", "Prepare standalone manuscript integration package", "Done", "PAPER_INTEGRATION_R4 prepared; locked main manuscript unchanged."),
    ("PF-GRP-CONSTRUCTIVE-R4-TOWER", "Construct two inequivalent towers", "Deferred", "Begin only after A1 freeze."),
]

common = {
    "Phase": "1 Parameter Freeze",
    "Category": "Finite quotient",
    "Priority": "P0",
    "Task Type": "Post-construction R4",
    "Exact Manuscript Contract": "Revision-4 post-construction closure directive",
    "Required Inputs": "CAND-R4-0005 mathematical root and typed downstream contracts",
    "Deliverable": "Content-addressed certificate/status artifact",
    "Acceptance Criteria": "Exact certificate and independent replay at the declared scope",
    "Forbidden Simplification / Failure Mode": "No degree-search restart; no invented resource value; no A1 promotion without Resource PASS",
    "Dependencies": "CAND-R4-0005-H-MATH-V1",
    "Owner": "Math/Code/Numerics",
    "Freeze Gate?": "Yes",
    "Numerical?": "No",
    "Source Location": "CONSTRUCTIVE_EXECUTION_R4",
}

wb = load_workbook(SOURCE, data_only=False, read_only=False)
ws = wb["Master Tasks"]
for offset, (task_id, task, status, notes) in enumerate(task_specs, start=1):
    row_no = ORIGINAL_LAST_ROW + offset
    record = dict(common)
    record.update({"Task ID": task_id, "Task": task, "Status": status, "Notes": notes})
    ws.row_dimensions[row_no].height = ws.row_dimensions[ORIGINAL_LAST_ROW].height
    for col, header in enumerate(headers, start=1):
        template = ws.cell(ORIGINAL_LAST_ROW, col)
        cell = ws.cell(row_no, col, record.get(header, ""))
        if template.has_style:
            cell._style = copy.copy(template._style)
        cell.font = copy.copy(template.font)
        cell.fill = copy.copy(template.fill)
        cell.border = copy.copy(template.border)
        cell.alignment = copy.copy(template.alignment)
        cell.protection = copy.copy(template.protection)
        cell.number_format = template.number_format

new_last_row = ORIGINAL_LAST_ROW + len(task_specs)
ws.tables["MasterTasksTable"].ref = f"A1:R{new_last_row}"

status_name = "R4 Post-Construction"
if status_name in wb.sheetnames:
    del wb[status_name]
status_ws = wb.create_sheet(status_name)
status_ws.append(["Field", "Consolidated value", "Evidence"])
with PREVIEW.open("r", encoding="utf-8", newline="") as handle:
    preview_rows = list(csv.DictReader(handle, delimiter="\t"))
for row in preview_rows:
    status_ws.append([row["Field"], row["Proposed value"], row["Evidence"]])
status_ws.append(["H_MATH", H_MATH, "17_FINAL_FREEZE/CAND-R4-0005_ROOT_CERTIFICATE.json"])
status_ws.append(["Workbook update policy", "SINGLE CONSOLIDATED VERSIONED UPDATE", "Frozen source remains immutable"])
for cell in status_ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center")
status_ws.freeze_panes = "A2"
status_ws.auto_filter.ref = f"A1:C{status_ws.max_row}"
status_ws.column_dimensions["A"].width = 38
status_ws.column_dimensions["B"].width = 72
status_ws.column_dimensions["C"].width = 76
for row in status_ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

if getattr(wb, "calculation", None) is None:
    wb.calculation = CalcProperties()
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(OUTPUT)

reloaded = load_workbook(OUTPUT, data_only=False, read_only=False)
rws = reloaded["Master Tasks"]
assert sha256(SOURCE) == source_hash_before
assert fingerprint(rws, ORIGINAL_LAST_ROW) == source_prefix
assert rws.max_row == new_last_row
assert rws.tables["MasterTasksTable"].ref == f"A1:R{new_last_row}"
assert [rws.cell(ORIGINAL_LAST_ROW + i, 1).value for i in range(1, len(task_specs) + 1)] == [row[0] for row in task_specs]
assert status_name in reloaded.sheetnames

source_formula_count = sum(1 for sheet in source_wb for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))
output_formula_count = sum(1 for sheet in reloaded for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("="))
assert source_formula_count == output_formula_count

audit = {
    "schema_version": "1.0",
    "classification": "PASS_SINGLE_CONSOLIDATED_VERSIONED_UPDATE",
    "source_sha256_before": source_hash_before,
    "source_sha256_after": sha256(SOURCE),
    "frozen_source_mutated": False,
    "output": "20_POSTCONSTRUCTION/Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan_POST_CONSTRUCTION_FINAL.xlsx",
    "output_sha256": sha256(OUTPUT),
    "original_region_preserved": True,
    "original_region_fingerprint": source_prefix,
    "task_rows_appended": len(task_specs),
    "new_master_task_last_row": new_last_row,
    "status_sheet": status_name,
    "status_fields": status_ws.max_row - 1,
    "old_degree_search_rows_preserved": True,
    "formula_count_source": source_formula_count,
    "formula_count_output": output_formula_count,
    "recalculate_on_open": True,
    "H_MATH": H_MATH,
}
AUDIT.write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(json.dumps({
    "classification": audit["classification"],
    "output_sha256": audit["output_sha256"],
    "task_rows_appended": len(task_specs),
    "status_fields": audit["status_fields"],
}, indent=2))
