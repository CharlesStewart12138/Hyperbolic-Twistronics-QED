"""Extract resource-relevant rows from the frozen authoritative workbook."""
from __future__ import annotations
import csv,hashlib,json
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2];R4=ROOT/"CONSTRUCTIVE_EXECUTION_R4"
BOOK=R4/"00_FROZEN_INPUTS"/"Hyperbolic_Bilayer_Parameter_Freeze_and_Exact_Code_Plan.xlsx"
ROWS=R4/"13_RESOURCE_GATE"/"FROZEN_WORKBOOK_RESOURCE_ROWS.tsv";AUDIT=R4/"13_RESOURCE_GATE"/"FROZEN_WORKBOOK_RESOURCE_AUDIT.json"
TERMS=("memory","resource","matvec","krylov","solver","wall","order","hilbert","dimension","degree","budget","rss","ram","storage","grid","probe","tolerance","precision","quotient","eigens","lanczos","arnoldi")
def main():
 wb=load_workbook(BOOK,data_only=False,read_only=False);rows=[];sheets={}
 for sheet in wb:
  sheets[sheet.title]={"max_row":sheet.max_row,"max_column":sheet.max_column,"range":sheet.calculate_dimension()}
  for cells in sheet.iter_rows():
   values=[cell.value for cell in cells];text=" | ".join("" if v is None else str(v) for v in values)
   matches=sorted({term for term in TERMS if term in text.lower()})
   if matches:rows.append({"sheet":sheet.title,"row":str(cells[0].row),"matched_terms":",".join(matches),"values":text})
 with ROWS.open("w",encoding="utf-8",newline="") as h:
  w=csv.DictWriter(h,delimiter="\t",fieldnames=["sheet","row","matched_terms","values"],lineterminator="\n");w.writeheader();w.writerows(rows)
 AUDIT.write_text(json.dumps({"schema_version":"1.0","workbook_sha256":hashlib.sha256(BOOK.read_bytes()).hexdigest().upper(),"sheets":sheets,
  "resource_relevant_rows":len(rows),"extraction_terms":TERMS,"output_rows_sha256":hashlib.sha256(ROWS.read_bytes()).hexdigest().upper(),
  "mutation_performed":False,"classification":"READ_ONLY_RESOURCE_EXTRACTION"},indent=2)+"\n",encoding="utf-8")
if __name__=="__main__":main()
