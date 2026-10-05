"""Public-dependency alternative to the original workbook catalog builder."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "02_WORKING" / "NUMERICAL_FIGURE_CATALOG_SOURCE.tsv"
QA = ROOT / "03_QA" / "NUMERICAL_FIGURE_QA.json"
OUTPUT = ROOT / "FINAL_NUMERICAL_FIGURE_CATALOG.xlsx"


def main() -> None:
    with SOURCE.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    qa = json.loads(QA.read_text(encoding="utf-8"))
    if len(rows) != 40 or qa.get("figure_count") != 40 or qa.get("status") != "PASS":
        raise RuntimeError("catalog source or QA payload is not a 40-figure PASS set")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Hyperbolic Bilayer — Numerical Figure Catalog"])
    summary.append([])
    summary.append(["Metric", "Value"])
    summary.append(["Total quantitative figures", len(rows)])
    summary.append(["QA status", qa["status"]])
    summary.append(["Schematic figures", 0])

    catalog = workbook.create_sheet("Figure Catalog")
    headers = list(rows[0])
    catalog.append(headers)
    for row in rows:
        catalog.append([row.get(header, "") for header in headers])

    deep = "4A251A"
    pale = "F8EEE7"
    for sheet in (summary, catalog):
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=deep)
            cell.font = Font(name="Times New Roman", color="FFFFFF", bold=True)
            cell.alignment = Alignment(vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Times New Roman", size=10)
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=pale)
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                55,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )

    workbook.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "rows": len(rows)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
